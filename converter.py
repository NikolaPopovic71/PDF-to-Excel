"""Pretvaranje PDF-a 'Medalists by event' (Splash Meet Manager) u Excel."""
import re
from collections import defaultdict

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_RE = re.compile(r"^(M/M|Z/W|Ž/W),\s+.+\s+\S+$")
TIME = r"(\d{1,2}:\d{2}\.\d{2}|\d{1,2}\.\d{2})"
PERSON_RE = re.compile(rf"^(\d+)\.\s+(.+?),\s*(.+?)\s+(\d{{2,4}})\s+([A-ZČĆŠŽĐ]{{2,4}})\s+{TIME}$")
RELAY_RE = re.compile(rf"^(\d+)\.\s+(.*?)\s*([A-ZČĆŠŽĐ]{{2,4}})\s+{TIME}$")
SKIP_RE = re.compile(r"^(Serbia OPEN|Beograd|Medalists|www\.|Splash)", re.I)


def title_case(s: str) -> str:
    """'DRAGANOV C' -> 'Draganov C', 'petrovic-jovic' -> 'Petrovic-Jovic'."""
    return re.sub(r"[^\s\-']+", lambda m: m.group(0)[0].upper() + m.group(0)[1:].lower(), s.strip())


def full_name(prezime: str, ime: str) -> str:
    """'DRAGANOV C', 'Stefanija' -> 'Stefanija C Draganov' (srednje slovo između imena i prezimena)."""
    tokens = prezime.split()
    srednje = [t for t in tokens if len(t.rstrip(".")) == 1]
    prez = [t for t in tokens if len(t.rstrip(".")) != 1]
    delovi = [title_case(ime)] + [s.upper() for s in srednje] + [title_case(" ".join(prez))]
    return " ".join(d for d in delovi if d)


def _lines(words):
    rows = defaultdict(list)
    for w in words:
        rows[round(w["top"] / 3)].append(w)
    for key in sorted(rows):
        ws = sorted(rows[key], key=lambda w: w["x0"])
        yield ws[0]["top"], " ".join(w["text"] for w in ws)


def extract_events(pdf_path):
    """Vraća listu disciplina: {'title': str, 'rows': [(rb, ime_prezime, godiste, klub, vreme)]}."""
    blocks = []  # (page, y, side, event)
    with pdfplumber.open(pdf_path) as pdf:
        for pno, page in enumerate(pdf.pages):
            mid = page.width / 2
            words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            for side in (0, 1):
                half = [w for w in words if (w["x0"] < mid) == (side == 0)]
                current = None
                for y, line in _lines(half):
                    if SKIP_RE.match(line):
                        continue
                    if HEADER_RE.match(line):
                        parts = line.rsplit(" ", 1)  # odvajamo "Open" za razmak
                        current = {"title": f"{parts[0]}     {parts[1]}", "rows": []}
                        blocks.append((pno, round(y), side, current))
                        continue
                    if current is None:
                        continue
                    m = PERSON_RE.match(line)
                    if m:
                        rb, prez, ime, god, klub, vreme = m.groups()
                        current["rows"].append((rb, full_name(prez, ime), god, klub, vreme))
                        continue
                    m = RELAY_RE.match(line)
                    if m:
                        rb, naziv, klub, vreme = m.groups()
                        current["rows"].append((rb, naziv.strip(), "", klub, vreme))
    blocks.sort(key=lambda b: (b[0], b[1], b[2]))
    return [b[3] for b in blocks]


def write_excel(events, xlsx_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rezultati"
    headers = ["R.br.", "Ime i prezime", "Godište", "Klub", "Vreme"]
    widths = [8, 38, 10, 10, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="1F4E78")
    head_fill = PatternFill("solid", fgColor="D9E1F2")
    center = Alignment(horizontal="center", vertical="center")

    r = 1
    for ev in events:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        c = ws.cell(r, 1, ev["title"])
        c.font = Font(bold=True, color="FFFFFF", size=12)
        c.fill = title_fill
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 20
        r += 1
        for col, h in enumerate(headers, 1):
            c = ws.cell(r, col, h)
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = center
            c.border = border
        r += 1
        for row in ev["rows"]:
            rb, ime, god, klub, vreme = row
            vals = [int(rb), ime, god, klub, vreme]
            for col, v in enumerate(vals, 1):
                c = ws.cell(r, col, v)
                c.border = border
                if col in (3, 5):
                    c.number_format = "@"  # tekst
                c.alignment = Alignment(horizontal="left" if col == 2 else "center")
            r += 1
        r += 1  # prazan red između disciplina
    wb.save(xlsx_path)


def convert(pdf_path, xlsx_path):
    events = extract_events(pdf_path)
    if not events:
        raise ValueError("U PDF-u nisu pronađene discipline. Da li je ovo 'Medalists by event' izveštaj?")
    write_excel(events, xlsx_path)
    return len(events), sum(len(e["rows"]) for e in events)
